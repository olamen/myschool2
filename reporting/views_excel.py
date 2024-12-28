import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from .models import NoteDevoir, Devoir, Classe, Subject, Trimestre, SessionYearModel

def export_notes_to_excel(request):
    # Récupérer les paramètres de filtre
    classe_id = request.GET.get("classe")
    devoir_id = request.GET.get("devoir")
    trimestre_id = request.GET.get("trimestre")
    subject_id = request.GET.get("subject")
    session_year_id = request.GET.get("session_year")

    # Filtrer les notes
    notes = NoteDevoir.objects.filter(
        student__student_class_id=classe_id,
        trimestre_id=trimestre_id,
        subject_id=subject_id,
        sessionyear_id=session_year_id,
    )

    if devoir_id:
        notes = notes.filter(devoir_id=devoir_id)
        devoir = get_object_or_404(Devoir, id=devoir_id)
    else:
        devoir = None

    # Création du fichier Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notes"

    # Définir le style des bordures
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Ajouter l'entête
    ws.merge_cells("A1:C1")
    ws["A1"] = "Nom de l'école en français"
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws["A1"].font = Font(size=14, bold=True)

    ws.merge_cells("D1:F1")
    ws["D1"] = "اسم المدرسة بالعربية"
    ws["D1"].alignment = Alignment(horizontal="right", vertical="center")
    ws["D1"].font = Font(size=14, bold=True)

    # Titre
    ws.merge_cells("A3:F3")
    ws["A3"] = f"Les Notes du Devoir : {devoir.name}" if devoir else "Les Notes des Devoirs"
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A3"].font = Font(size=16, bold=True)

    # Ajouter les colonnes
    headers = ["#", "Nom de l'Étudiant", "Note"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Ajouter les données des étudiants
    for index, note in enumerate(notes, start=1):
        ws.cell(row=6 + index, column=1, value=index).border = thin_border
        ws.cell(row=6 + index, column=2, value=f"{note.student.first_name} {note.student.last_name}").border = thin_border
        ws.cell(row=6 + index, column=3, value=note.score).border = thin_border

    # Ajuster la largeur des colonnes
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 10

    # Générer le fichier Excel
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=notes_devoir.xlsx"
    wb.save(response)
    return response